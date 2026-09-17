"""Periodic cumulative AE/SAE safety summary report — Phase 3 of the AE/SAE
project (see project memory). A site-wise, date-ranged line-listing built
from the two existing AE data sources: `sae_reports` (the CIOMS-shaped SAE
detail — causality/seriousness/outcome) and `adverse_events` (the broader
per-baby AE register — every clinician-accepted AE, shallower fields).
Formatted as .docx, matching sae_report.py's conventions.

The caller (routers/dashboard.py) assembles all data from the database;
this module is DB-free, same discipline as ae_reference.py/sae_report.py.
"""

from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT

import sae_config as cfg

_BLANK = "—"


def _v(x):
    if x is None:
        return _BLANK
    s = str(x).strip()
    return s or _BLANK


def _doc():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    for section in d.sections:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
    return d


def _h(d, text, size=13, space_before=12):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    return p


def _para(d, text, bold=False, italic=False, size=10):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    return p


def _site_display(site):
    return cfg.SITES.get(site, {}).get("display", site)


def _listing_table(d, columns, rows):
    t = d.add_table(rows=1, cols=len(columns))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, label in enumerate(columns):
        r = hdr[i].paragraphs[0].add_run(label)
        r.bold = True
        r.font.size = Pt(8.5)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].paragraphs[0].add_run(_v(val)).font.size = Pt(8.5)


def _site_section(d, site, rows, columns, row_mapper, empty_note):
    _para(d, _site_display(site), bold=True, size=10.5)
    if not rows:
        _para(d, empty_note, italic=True, size=9)
        d.add_paragraph()
        return
    _listing_table(d, columns, [row_mapper(r) for r in rows])
    d.add_paragraph()


SAE_COLUMNS = (
    "Enrolment ID", "Onset", "Diagnosis", "Seriousness criteria",
    "Severity", "Causality", "Action taken", "Outcome",
)
AE_COLUMNS = ("Enrolment ID", "Onset", "AE description", "Grade", "Converted to SAE")


def _sae_row(r):
    return (
        r.get("enrollment_id"), r.get("onset"), r.get("diagnosis"),
        r.get("seriousness"), r.get("severity_label"), r.get("causality"),
        r.get("action_taken"), r.get("outcome"),
    )


def _ae_row(r):
    return (
        r.get("enrollment_id"), r.get("onset"), r.get("description"),
        r.get("grade_label"), r.get("converted_to_sae"),
    )


def build_summary_report_docx(ctx) -> bytes:
    ctx = ctx or {}
    d = _doc()

    title = d.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("Cumulative Adverse Event & SAE Safety Summary")
    tr.bold = True
    tr.font.size = Pt(15)
    sub = d.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(cfg.TRIAL["protocol_title"])
    sr.italic = True
    sr.font.size = Pt(9.5)

    _para(
        d,
        f"Period covered: {ctx.get('date_from') or 'trial start'} to "
        f"{ctx.get('date_to') or 'present'}",
        bold=True,
    )

    sites = ctx.get("sites") or list(cfg.SITES.keys())
    counts = ctx.get("counts") or {}
    overall = counts.get("__overall__", {"n_ae": 0, "n_sae": 0})

    _h(d, "Summary counts by site")
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for i, label in enumerate(("Site", "All adverse events", "Serious adverse events (SAEs)")):
        r = t.rows[0].cells[i].paragraphs[0].add_run(label)
        r.bold = True
        r.font.size = Pt(9.5)
    row = t.add_row().cells
    row[0].paragraphs[0].add_run("Overall").bold = True
    row[1].paragraphs[0].add_run(str(overall.get("n_ae", 0))).font.size = Pt(9.5)
    row[2].paragraphs[0].add_run(str(overall.get("n_sae", 0))).font.size = Pt(9.5)
    for s in sites:
        c = counts.get(s, {"n_ae": 0, "n_sae": 0})
        row = t.add_row().cells
        row[0].paragraphs[0].add_run(_site_display(s)).font.size = Pt(9.5)
        row[1].paragraphs[0].add_run(str(c.get("n_ae", 0))).font.size = Pt(9.5)
        row[2].paragraphs[0].add_run(str(c.get("n_sae", 0))).font.size = Pt(9.5)

    _h(d, "Section A — Serious Adverse Events (line listing)")
    _para(
        d,
        "Source: SAE Reporting Form records (sae_reports). Severity per the INC NAESS "
        "5-grade scale; causality/outcome as entered by the site investigator. This is "
        "the CIOMS-shaped listing (causality/seriousness/outcome captured).",
        italic=True, size=9,
    )
    sae_by_site = ctx.get("sae_by_site") or {}
    for s in sites:
        _site_section(
            d, s, sae_by_site.get(s, []), SAE_COLUMNS, _sae_row,
            "No SAEs recorded for this site in the selected period.",
        )

    _h(d, "Section B — All Adverse Events (line listing)")
    _para(
        d,
        "Source: the per-baby Adverse Events register (adverse_events) — every AE a "
        "clinician has accepted onto the register, including non-serious events. "
        "Broader but shallower than Section A: no causality/seriousness/outcome is "
        "captured for these rows.",
        italic=True, size=9,
    )
    ae_by_site = ctx.get("ae_by_site") or {}
    for s in sites:
        _site_section(
            d, s, ae_by_site.get(s, []), AE_COLUMNS, _ae_row,
            "No adverse events recorded for this site in the selected period.",
        )

    foot = d.add_paragraph()
    fr = foot.add_run(
        f"\nGenerated by the PORTAL Trial data system on "
        f"{datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
        + (f" · {ctx.get('generated_by')}" if ctx.get("generated_by") else "")
        + ". Tabulated summary for periodic IEC/DSMC review — submit as a hard copy per "
        "site convention (CIOMS practice for tabulated summaries)."
    )
    fr.italic = True
    fr.font.size = Pt(8)

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
